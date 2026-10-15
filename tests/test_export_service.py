from datetime import date
from decimal import Decimal
import openpyxl
import pytest

from models.gasto import Gasto
from repository.categoria_repo import CategoriaRepo
from repository.gasto_repo import GastoRepo
from repository.miembro_repo import MiembroRepo
from services.export_service import ExportService
from services.resumen_service import ResumenService


def test_export_service_excel_y_csv(
    tmp_path,
    gasto_repo: GastoRepo,
    categoria_repo: CategoriaRepo,
    miembro_repo: MiembroRepo,
    resumen_service: ResumenService,
):
    super_cat = categoria_repo.obtener_por_nombre("Supermercado")
    serv_cat = categoria_repo.obtener_por_nombre("Servicios")
    miembro = miembro_repo.crear(miembro_repo.listar()[0] if miembro_repo.listar() else type("M", (), {"id": None, "nombre": "Santiago"})())

    # Cargar gastos de prueba
    gasto_repo.crear(
        Gasto(
            fecha=date(2026, 10, 5),
            monto=Decimal("1500.50"),
            categoria_id=super_cat.id,
            descripcion="Supermercado mensual",
            medio_pago="Débito",
            miembro_id=miembro.id,
        )
    )
    gasto_repo.crear(
        Gasto(
            fecha=date(2026, 10, 15),
            monto=Decimal("800.00"),
            categoria_id=serv_cat.id,
            descripcion="Internet",
            medio_pago="Transferencia",
        )
    )

    export_service = ExportService(
        gasto_repo=gasto_repo,
        categoria_repo=categoria_repo,
        miembro_repo=miembro_repo,
        resumen_service=resumen_service,
    )

    # 1. Probar exportación a Excel
    ruta_excel = tmp_path / "reporte_gastos.xlsx"
    export_service.exportar_a_excel(ruta_excel, 2026, 10)
    assert ruta_excel.exists()

    # Cargar y verificar hojas con openpyxl
    wb = openpyxl.load_workbook(ruta_excel)
    assert "Gastos" in wb.sheetnames
    assert "Resumen por Categoría" in wb.sheetnames

    ws_gastos = wb["Gastos"]
    assert ws_gastos.cell(row=1, column=1).value == "Fecha"
    assert ws_gastos.cell(row=2, column=2).value == "Supermercado"
    assert ws_gastos.cell(row=2, column=4).value == 1500.50

    ws_resumen = wb["Resumen por Categoría"]
    assert ws_resumen.cell(row=1, column=1).value == "Categoría"
    assert ws_resumen.cell(row=2, column=1).value == "Supermercado"
    assert ws_resumen.cell(row=2, column=2).value == 1500.50

    # 2. Probar exportación a CSV
    ruta_csv = tmp_path / "reporte_gastos.csv"
    export_service.exportar_a_csv(ruta_csv, 2026, 10)
    assert ruta_csv.exists()

    with open(ruta_csv, "r", encoding="utf-8-sig") as f:
        contenido = f.read()
        assert "Fecha;Categoría;Descripción;Monto" in contenido
        assert "Supermercado mensual" in contenido
        assert "Internet" in contenido
