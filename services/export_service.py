import csv
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import formatear_moneda
from repository.categoria_repo import CategoriaRepo
from repository.gasto_repo import GastoRepo
from repository.miembro_repo import MiembroRepo
from services.resumen_service import ResumenService


class ExportService:
    """Servicio de exportación de gastos y resúmenes a formatos Excel (.xlsx) y CSV."""

    def __init__(
        self,
        gasto_repo: GastoRepo,
        categoria_repo: CategoriaRepo,
        miembro_repo: MiembroRepo,
        resumen_service: ResumenService,
    ):
        self.gasto_repo = gasto_repo
        self.categoria_repo = categoria_repo
        self.miembro_repo = miembro_repo
        self.resumen_service = resumen_service

    def exportar_a_excel(
        self, ruta_archivo: Path | str, anio: int, mes: int
    ) -> Path:
        """Genera un archivo Excel con dos hojas: 'Gastos' y 'Resumen por Categoría'."""
        wb = Workbook()

        # Estilos para cabeceras
        font_cabecera = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill_cabecera = PatternFill(
            start_color="2C3E50", end_color="2C3E50", fill_type="solid"
        )
        align_centro = Alignment(horizontal="center", vertical="center")
        align_derecha = Alignment(horizontal="right", vertical="center")
        borde_fino = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0"),
        )
        font_total = Font(name="Calibri", size=11, bold=True)

        # ----------------- HOJA 1: GASTOS DETALLADOS -----------------
        ws_gastos = wb.active
        ws_gastos.title = "Gastos"

        cabeceras_gastos = [
            "Fecha",
            "Categoría",
            "Descripción",
            "Monto ($)",
            "Medio de Pago",
            "Pagó",
        ]
        ws_gastos.append(cabeceras_gastos)

        for col_idx in range(1, len(cabeceras_gastos) + 1):
            celda = ws_gastos.cell(row=1, column=col_idx)
            celda.font = font_cabecera
            celda.fill = fill_cabecera
            celda.alignment = align_centro

        gastos = self.gasto_repo.listar_por_mes(anio, mes)
        cats = {c.id: c for c in self.categoria_repo.listar(solo_activas=False)}
        miembros = {m.id: m for m in self.miembro_repo.listar()}

        fila_actual = 2
        for g in gastos:
            cat_nom = cats[g.categoria_id].nombre if g.categoria_id in cats else "General"
            miembro_nom = (
                miembros[g.miembro_id].nombre
                if g.miembro_id and g.miembro_id in miembros
                else ""
            )

            ws_gastos.append(
                [
                    g.fecha.strftime("%d/%m/%Y"),
                    cat_nom,
                    g.descripcion,
                    float(g.monto),
                    g.medio_pago,
                    miembro_nom,
                ]
            )

            # Formato de celda de monto como moneda
            celda_monto = ws_gastos.cell(row=fila_actual, column=4)
            celda_monto.number_format = "$ #,##0.00"
            celda_monto.alignment = align_derecha

            # Formato fecha centrada
            ws_gastos.cell(row=fila_actual, column=1).alignment = align_centro

            for c in range(1, len(cabeceras_gastos) + 1):
                ws_gastos.cell(row=fila_actual, column=c).border = borde_fino

            fila_actual += 1

        # Fila de Total en Hoja Gastos
        if fila_actual > 2:
            celda_lbl = ws_gastos.cell(row=fila_actual, column=3, value="Total del Mes:")
            celda_lbl.font = font_total
            celda_lbl.alignment = align_derecha

            celda_sum = ws_gastos.cell(
                row=fila_actual,
                column=4,
                value=f"=SUM(D2:D{fila_actual-1})",
            )
            celda_sum.font = font_total
            celda_sum.number_format = "$ #,##0.00"
            celda_sum.alignment = align_derecha

        # ----------------- HOJA 2: RESUMEN POR CATEGORÍA -----------------
        ws_resumen = wb.create_sheet(title="Resumen por Categoría")
        cabeceras_resumen = ["Categoría", "Total Gastado ($)", "% del Total"]
        ws_resumen.append(cabeceras_resumen)

        for col_idx in range(1, len(cabeceras_resumen) + 1):
            celda = ws_resumen.cell(row=1, column=col_idx)
            celda.font = font_cabecera
            celda.fill = PatternFill(
                start_color="16A085", end_color="16A085", fill_type="solid"
            )
            celda.alignment = align_centro

        totales_cat = self.resumen_service.totales_por_categoria(anio, mes)
        fila_res = 2
        for r in totales_cat:
            ws_resumen.append(
                [r["nombre"], float(r["total"]), float(r["porcentaje"]) / 100.0]
            )
            # Monto moneda
            c_monto = ws_resumen.cell(row=fila_res, column=2)
            c_monto.number_format = "$ #,##0.00"
            c_monto.alignment = align_derecha

            # Porcentaje
            c_pct = ws_resumen.cell(row=fila_res, column=3)
            c_pct.number_format = "0.0%"
            c_pct.alignment = align_derecha

            for c in range(1, len(cabeceras_resumen) + 1):
                ws_resumen.cell(row=fila_res, column=c).border = borde_fino

            fila_res += 1

        # Autoajustar ancho de columnas en ambas hojas
        for ws in (ws_gastos, ws_resumen):
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        ruta = Path(ruta_archivo)
        wb.save(ruta)
        return ruta

    def exportar_a_csv(
        self, ruta_archivo: Path | str, anio: int, mes: int
    ) -> Path:
        """Exporta los gastos a CSV compatible con UTF-8 para Excel."""
        gastos = self.gasto_repo.listar_por_mes(anio, mes)
        cats = {c.id: c for c in self.categoria_repo.listar(solo_activas=False)}
        miembros = {m.id: m for m in self.miembro_repo.listar()}

        ruta = Path(ruta_archivo)
        with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerow(
                [
                    "Fecha",
                    "Categoría",
                    "Descripción",
                    "Monto",
                    "Medio de Pago",
                    "Pagó",
                ]
            )
            for g in gastos:
                cat_nom = cats[g.categoria_id].nombre if g.categoria_id in cats else ""
                miembro_nom = (
                    miembros[g.miembro_id].nombre
                    if g.miembro_id and g.miembro_id in miembros
                    else ""
                )
                writer.writerow(
                    [
                        g.fecha.strftime("%d/%m/%Y"),
                        cat_nom,
                        g.descripcion,
                        str(g.monto).replace(".", ","),
                        g.medio_pago,
                        miembro_nom,
                    ]
                )

        return ruta
